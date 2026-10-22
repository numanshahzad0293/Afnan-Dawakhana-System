import os, sys, sqlite3, tempfile, webbrowser, csv, html
from datetime import datetime, date
import tkinter as tk
from tkinter import ttk, messagebox, filedialog

APP_NAME = "Afnan Shahzad Dawakhana"
PHONE = "0333-7676293"
ADDRESS = "Near Mand Barfi Shop, Tehsil Office Road, Gojra"
HOURS = "Every Sunday: 09:00 AM - 04:00 PM"
DB_NAME = "afnan_dawakhana.db"

def app_data_dir():
    root = os.environ.get("LOCALAPPDATA") or os.path.expanduser("~")
    path = os.path.join(root, "AfnanDawakhana")
    os.makedirs(path, exist_ok=True)
    return path

DB_PATH = os.path.join(app_data_dir(), DB_NAME)

def today_key():
    return date.today().isoformat()

def date_label(value=None):
    value = value or date.today().isoformat()
    try:
        return datetime.fromisoformat(value).strftime("%d %b %Y")
    except Exception:
        return value

class DawakhanaApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Afnan Shahzad Dawakhana - Patient & Token System")
        self.root.geometry("1180x780")
        self.root.minsize(960, 650)
        self.conn = sqlite3.connect(DB_PATH)
        self.conn.row_factory = sqlite3.Row
        self.init_db()
        self.build_ui()
        self.refresh_records()
        self.update_token_display()

    def init_db(self):
        self.conn.execute("""CREATE TABLE IF NOT EXISTS patients (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            token_week TEXT NOT NULL,
            token_no INTEGER NOT NULL,
            patient_name TEXT NOT NULL,
            mobile TEXT,
            address TEXT,
            disease TEXT,
            created_at TEXT NOT NULL
        )""")
        self.conn.execute("CREATE INDEX IF NOT EXISTS idx_patient_name ON patients(patient_name)")
        self.conn.execute("CREATE INDEX IF NOT EXISTS idx_patient_mobile ON patients(mobile)")
        self.conn.execute("""CREATE TABLE IF NOT EXISTS settings (
            key TEXT PRIMARY KEY, value TEXT NOT NULL
        )""")
        self.conn.execute("""CREATE TABLE IF NOT EXISTS token_counters (
            token_date TEXT PRIMARY KEY, last_token INTEGER NOT NULL DEFAULT 0
        )""")
        defaults = {"clinic_name": APP_NAME, "clinic_address": ADDRESS, "clinic_phone": PHONE, "clinic_hours": HOURS}
        for key, value in defaults.items():
            self.conn.execute("INSERT OR IGNORE INTO settings(key,value) VALUES(?,?)", (key, value))
        self.conn.commit()

    def setting(self, key):
        row = self.conn.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
        return row["value"] if row else ""

    def set_setting(self, key, value):
        self.conn.execute("INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (key, value))
        self.conn.commit()

    def build_ui(self):
        style = ttk.Style()
        try: style.theme_use("clam")
        except tk.TclError: pass

        header = ttk.Frame(self.root, padding=(16, 10))
        header.pack(fill="x")
        self.header_name = ttk.Label(header, text=self.setting("clinic_name"), font=("Segoe UI", 19, "bold"))
        self.header_name.pack(anchor="w")
        self.header_details = ttk.Label(header, text=f"{self.setting('clinic_address')}   |   {self.setting('clinic_hours')}", font=("Segoe UI", 10))
        self.header_details.pack(anchor="w", pady=(3, 0))
        ttk.Button(header, text="Settings", command=self.open_settings).pack(anchor="e", pady=(0, 2))

        main = ttk.Frame(self.root, padding=(16, 2, 16, 12)); main.pack(fill="both", expand=True)
        form = ttk.LabelFrame(main, text="New Patient / Token", padding=10); form.pack(fill="x")
        self.name_var = tk.StringVar(); self.mobile_var = tk.StringVar(); self.address_var = tk.StringVar(); self.disease_var = tk.StringVar()
        self.token_var = tk.StringVar(value="001")
        ttk.Label(form, text="Token Number").grid(row=0,column=0,sticky="w",padx=5,pady=4)
        ttk.Label(form, textvariable=self.token_var, font=("Segoe UI",20,"bold")).grid(row=1,column=0,sticky="w",padx=5,pady=(0,6))
        ttk.Label(form,text="Patient Name *").grid(row=0,column=1,sticky="w",padx=5,pady=4)
        ttk.Entry(form,textvariable=self.name_var).grid(row=1,column=1,sticky="ew",padx=5,pady=(0,6))
        ttk.Label(form,text="Mobile Number").grid(row=0,column=2,sticky="w",padx=5,pady=4)
        ttk.Entry(form,textvariable=self.mobile_var).grid(row=1,column=2,sticky="ew",padx=5,pady=(0,6))
        ttk.Label(form,text="Patient Address").grid(row=2,column=0,sticky="w",padx=5,pady=4)
        ttk.Entry(form,textvariable=self.address_var).grid(row=3,column=0,columnspan=2,sticky="ew",padx=5,pady=(0,6))
        ttk.Label(form,text="Disease / Problem (internal only)").grid(row=2,column=2,sticky="w",padx=5,pady=4)
        ttk.Entry(form,textvariable=self.disease_var).grid(row=3,column=2,sticky="ew",padx=5,pady=(0,6))
        actions=ttk.Frame(form); actions.grid(row=4,column=0,columnspan=3,sticky="w",padx=5,pady=(4,0))
        ttk.Button(actions,text="Save Patient",command=self.save_patient).pack(side="left",padx=(0,7))
        ttk.Button(actions,text="Save & Print Slip",command=lambda:self.save_patient(True)).pack(side="left",padx=(0,7))
        ttk.Button(actions,text="Clear Form",command=self.clear_form).pack(side="left",padx=(0,7))
        ttk.Button(actions,text="Manual Token Reset",command=self.manual_reset).pack(side="left")
        for c in range(3): form.columnconfigure(c,weight=1)

        search=ttk.LabelFrame(main,text="Search / Filter Patient Records",padding=8); search.pack(fill="x",pady=(10,8))
        self.search_var=tk.StringVar()
        search_entry=ttk.Entry(search,textvariable=self.search_var)
        search_entry.pack(side="left",fill="x",expand=True,padx=(0,7))
        search_entry.bind("<Return>",lambda _e:self.refresh_records())
        ttk.Button(search,text="Search",command=self.refresh_records).pack(side="left",padx=(0,6))
        ttk.Button(search,text="Show All",command=self.show_all).pack(side="left",padx=(0,6))
        ttk.Label(search,text="From (YYYY-MM-DD)").pack(side="left",padx=(8,3))
        self.from_var=tk.StringVar(value=date.today().isoformat())
        ttk.Entry(search,textvariable=self.from_var,width=12).pack(side="left",padx=(0,5))
        ttk.Label(search,text="To").pack(side="left",padx=(0,3))
        self.to_var=tk.StringVar(value=date.today().isoformat())
        ttk.Entry(search,textvariable=self.to_var,width=12).pack(side="left",padx=(0,5))
        ttk.Button(search,text="Filter Dates",command=self.refresh_records).pack(side="left",padx=(0,6))
        ttk.Button(search,text="Export Excel",command=self.export_excel).pack(side="left")

        records=ttk.LabelFrame(main,text="Saved Patients (select one or more rows for printing)",padding=7); records.pack(fill="both",expand=True)
        cols=("token","name","mobile","address","disease","date")
        self.tree=ttk.Treeview(records,columns=cols,show="headings",height=12,selectmode="extended")
        headings={"token":"Token","name":"Patient Name","mobile":"Mobile","address":"Address","disease":"Disease / Problem","date":"Date"}
        widths={"token":65,"name":160,"mobile":120,"address":250,"disease":180,"date":110}
        for c in cols:
            self.tree.heading(c,text=headings[c]); self.tree.column(c,width=widths[c],anchor="w")
        ys=ttk.Scrollbar(records,orient="vertical",command=self.tree.yview); xs=ttk.Scrollbar(records,orient="horizontal",command=self.tree.xview)
        self.tree.configure(yscrollcommand=ys.set,xscrollcommand=xs.set)
        self.tree.grid(row=0,column=0,sticky="nsew"); ys.grid(row=0,column=1,sticky="ns"); xs.grid(row=1,column=0,sticky="ew")
        records.rowconfigure(0,weight=1); records.columnconfigure(0,weight=1)
        bottom=ttk.Frame(main); bottom.pack(fill="x",pady=(7,0))
        ttk.Label(bottom,text=f"Contact: {self.setting('clinic_phone')}   |   Patient data saved locally on this computer").pack(side="left")
        ttk.Button(bottom,text="Print Selected Slip(s)",command=self.print_selected).pack(side="right",padx=(6,0))
        ttk.Button(bottom,text="Print Preview",command=self.preview_selected).pack(side="right",padx=(6,0))

    def next_token(self):
        d=today_key()
        row=self.conn.execute("SELECT last_token FROM token_counters WHERE token_date=?",(d,)).fetchone()
        if row: return int(row["last_token"])+1
        row=self.conn.execute("SELECT COALESCE(MAX(token_no),0) n FROM patients WHERE date(created_at)=date(?)",(d,)).fetchone()
        return int(row["n"] or 0)+1

    def update_token_display(self): self.token_var.set(f"{self.next_token():03d}")

    def save_patient(self, print_after=False):
        name=self.name_var.get().strip()
        if not name:
            messagebox.showwarning("Patient Name Required","Please enter the patient's name."); return
        d=today_key(); token=self.next_token()
        self.conn.execute("INSERT INTO token_counters(token_date,last_token) VALUES(?,?) ON CONFLICT(token_date) DO UPDATE SET last_token=excluded.last_token",(d,token))
        created=datetime.now().isoformat(timespec="seconds")
        cur=self.conn.execute("INSERT INTO patients(token_week,token_no,patient_name,mobile,address,disease,created_at) VALUES(?,?,?,?,?,?,?)",(d,token,name,self.mobile_var.get().strip(),self.address_var.get().strip(),self.disease_var.get().strip(),created))
        self.conn.commit(); patient=self.conn.execute("SELECT * FROM patients WHERE id=?",(cur.lastrowid,)).fetchone()
        self.refresh_records(); self.clear_form()
        if print_after: self.print_batch([patient],preview=False)
        messagebox.showinfo("Saved",f"Patient saved successfully. Token number: {token:03d}")

    def clear_form(self):
        self.name_var.set(""); self.mobile_var.set(""); self.address_var.set(""); self.disease_var.set(""); self.update_token_display()

    def show_all(self):
        self.search_var.set(""); self.from_var.set(""); self.to_var.set(""); self.refresh_records()

    def refresh_records(self):
        for item in self.tree.get_children(): self.tree.delete(item)
        q=self.search_var.get().strip(); start=self.from_var.get().strip(); end=self.to_var.get().strip()
        sql="SELECT * FROM patients WHERE 1=1"; args=[]
        if q:
            like=f"%{q}%"; sql += " AND (patient_name LIKE ? OR mobile LIKE ? OR address LIKE ? OR disease LIKE ?)"; args += [like,like,like,like]
        try:
            if start:
                datetime.strptime(start,"%Y-%m-%d"); sql += " AND date(created_at) >= date(?)"; args.append(start)
            if end:
                datetime.strptime(end,"%Y-%m-%d"); sql += " AND date(created_at) <= date(?)"; args.append(end)
        except ValueError:
            messagebox.showwarning("Invalid Date","Please use YYYY-MM-DD date format."); return
        sql += " ORDER BY id DESC LIMIT 5000"
        rows=self.conn.execute(sql,args).fetchall()
        for r in rows:
            self.tree.insert("","end",iid=str(r["id"]),values=(f"{r['token_no']:03d}",r["patient_name"],r["mobile"] or "",r["address"] or "",r["disease"] or "",date_label(r["created_at"][:10])))
        self.update_token_display()

    def selected_rows(self):
        ids=self.tree.selection()
        if not ids:
            messagebox.showwarning("Select Patient","Please select one or more patient records first."); return []
        rows=[]
        for ident in ids:
            row=self.conn.execute("SELECT * FROM patients WHERE id=?",(int(ident),)).fetchone()
            if row: rows.append(row)
        return rows

    def print_selected(self):
        rows=self.selected_rows()
        if rows: self.print_batch(rows,preview=False)

    def preview_selected(self):
        rows=self.selected_rows()
        if rows: self.print_batch(rows,preview=True)

    @staticmethod
    def esc(value): return html.escape(str(value or ""),quote=True)

    def print_batch(self, rows, preview=False):
        pages=[]
        for p in rows:
            token=f"{p['token_no']:03d}"; dt=date_label(p['created_at'][:10])
            pages.append(f'''<section class="slip"><div class="center clinic">{self.esc(self.setting('clinic_name'))}</div>
<div class="center addr">{self.esc(self.setting('clinic_address'))}</div><div class="dash"></div>
<div class="center hours">{self.esc(self.setting('clinic_hours'))}</div>
<div class="tokenbox"><div class="tokenlabel">TOKEN NO.</div><div class="token">{token}</div></div>
<div class="row"><b>Patient Name</b><div class="value">{self.esc(p['patient_name'])}</div></div>
<div class="row"><b>Mobile Number</b><div class="value">{self.esc(p['mobile'])}</div></div>
<div class="row"><b>Address</b><div class="value">{self.esc(p['address'])}</div></div>
<div class="dash"></div><div class="center footer">For more info</div><div class="center phone">{self.esc(self.setting('clinic_phone'))}</div><div class="center dt">{dt}</div></section>''')
        controls="" if not preview else '<div class="controls"><button onclick="window.print()">Print</button><button onclick="window.close()">Close Preview</button></div>'
        html_doc=f'''<!doctype html><html><head><meta charset="utf-8"><title>Token Print Preview</title><style>
*{{box-sizing:border-box}} body{{margin:0;font-family:Arial,sans-serif;color:#000;background:#eee}} .controls{{position:sticky;top:0;background:#ddd;padding:10px;text-align:center}} button{{padding:8px 18px;margin:0 6px}} .slip{{width:74mm;margin:8mm auto;background:#fff;padding:2mm 1mm;page-break-after:always;break-after:page;overflow-wrap:anywhere}} .slip:last-child{{page-break-after:auto;break-after:auto}} .center{{text-align:center}} .clinic{{font-size:16pt;font-weight:800;line-height:1.05}} .addr{{font-size:8.5pt;margin-top:2mm}} .dash{{border-top:1px dashed #000;margin:2mm 0}} .hours{{font-size:9pt;font-weight:bold}} .tokenbox{{border:2px solid #000;border-radius:3mm;text-align:center;padding:1.5mm;margin:2mm 0}} .tokenlabel{{font-size:12pt;font-weight:bold}} .token{{font-size:34pt;font-weight:900;line-height:1}} .row{{margin:2mm 0;font-size:9.5pt}} .value{{border-bottom:1px dotted #000;min-height:4mm;padding:0.5mm 0}} .footer{{font-size:8.5pt;margin-top:2mm}} .phone{{font-size:12pt;font-weight:900}} .dt{{font-size:8pt;margin-top:1mm}} @page{{size:80mm auto;margin:2mm}} @media print{{body{{background:#fff}} .controls{{display:none}} .slip{{width:auto;margin:0;padding:0;box-shadow:none}}}}
</style></head><body>{controls}{''.join(pages)}{'' if preview else '<script>window.onload=function(){window.print();};</script>'}</body></html>'''
        out=os.path.join(tempfile.gettempdir(),"afnan_token_preview.html" if preview else "afnan_token_print.html")
        with open(out,"w",encoding="utf-8") as f: f.write(html_doc)
        webbrowser.open("file:///"+out.replace("\\","/"))

    def manual_reset(self):
        if not messagebox.askyesno("Confirm Token Reset","Reset today's token counter to 000 so the next patient gets 001? All saved patient records will remain unchanged."):
            return
        self.conn.execute("INSERT INTO token_counters(token_date,last_token) VALUES(?,0) ON CONFLICT(token_date) DO UPDATE SET last_token=0",(today_key(),))
        self.conn.commit(); self.update_token_display()
        messagebox.showinfo("Token Reset","Today's token counter is reset. Existing patient records are preserved.")

    def export_excel(self):
        start=self.from_var.get().strip(); end=self.to_var.get().strip()
        try:
            if start: datetime.strptime(start,"%Y-%m-%d")
            if end: datetime.strptime(end,"%Y-%m-%d")
        except ValueError:
            messagebox.showwarning("Invalid Date","Please use YYYY-MM-DD date format."); return
        sql="SELECT * FROM patients WHERE 1=1"; args=[]
        if start: sql += " AND date(created_at)>=date(?)"; args.append(start)
        if end: sql += " AND date(created_at)<=date(?)"; args.append(end)
        sql += " ORDER BY created_at, token_no"
        rows=self.conn.execute(sql,args).fetchall()
        if not rows:
            messagebox.showinfo("No Records","No patient records found for the selected date range."); return
        path=filedialog.asksaveasfilename(title="Save patient report",defaultextension=".xlsx",filetypes=[("Excel Workbook","*.xlsx"),("CSV file","*.csv")],initialfile=f"Patient_Report_{start or 'all'}_to_{end or 'all'}.xlsx")
        if not path: return
        headers=["Token Number","Patient Name","Mobile Number","Address","Disease / Problem","Date","Time"]
        data=[[f"{r['token_no']:03d}",r['patient_name'],r['mobile'] or "",r['address'] or "",r['disease'] or "",r['created_at'][:10],r['created_at'][11:19]] for r in rows]
        try:
            if path.lower().endswith(".csv"):
                with open(path,"w",newline="",encoding="utf-8-sig") as f:
                    writer=csv.writer(f); writer.writerow(headers); writer.writerows(data)
            else:
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill, Alignment
                wb=Workbook(); ws=wb.active; ws.title="Patients"; ws.append(headers)
                for row in data: ws.append(row)
                for cell in ws[1]:
                    cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="1F4E78"); cell.alignment=Alignment(horizontal="center")
                for col,width in {"A":15,"B":26,"C":18,"D":48,"E":28,"F":16,"G":14}.items(): ws.column_dimensions[col].width=width
                ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions; wb.save(path)
            messagebox.showinfo("Export Complete",f"Exported {len(rows)} patient records to:\n{path}")
        except ImportError:
            messagebox.showerror("Excel Library Missing","Excel export needs openpyxl. Please rebuild the EXE with openpyxl included.")
        except Exception as e: messagebox.showerror("Export Failed",str(e))

    def open_settings(self):
        win=tk.Toplevel(self.root); win.title("Clinic Settings"); win.resizable(False,False); win.transient(self.root); win.grab_set()
        fields=[("Clinic Name","clinic_name"),("Clinic Address","clinic_address"),("Mobile Number","clinic_phone"),("Clinic Hours","clinic_hours")]
        vars_={}; frame=ttk.Frame(win,padding=16); frame.pack(fill="both",expand=True)
        for i,(label,key) in enumerate(fields):
            ttk.Label(frame,text=label).grid(row=i,column=0,sticky="w",padx=(0,10),pady=7)
            var=tk.StringVar(value=self.setting(key)); vars_[key]=var
            ttk.Entry(frame,textvariable=var,width=48).grid(row=i,column=1,sticky="ew",pady=7)
        def save():
            for key,var in vars_.items(): self.set_setting(key,var.get().strip())
            self.header_name.configure(text=self.setting("clinic_name"))
            self.header_details.configure(text=f"{self.setting('clinic_address')}   |   {self.setting('clinic_hours')}")
            messagebox.showinfo("Settings Saved","Clinic details saved and will appear on new printouts.",parent=win); win.destroy()
        buttons=ttk.Frame(frame); buttons.grid(row=len(fields),column=0,columnspan=2,sticky="e",pady=(10,0))
        ttk.Button(buttons,text="Save Settings",command=save).pack(side="left",padx=5); ttk.Button(buttons,text="Cancel",command=win.destroy).pack(side="left")

    def on_close(self): self.conn.close(); self.root.destroy()

def main():
    root=tk.Tk(); app=DawakhanaApp(root); root.protocol("WM_DELETE_WINDOW",app.on_close); root.mainloop()

if __name__ == "__main__": main()
